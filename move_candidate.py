#!/usr/bin/env python3
"""
Move one candidate's row between sheets of the Margin Dashboard workbook
(e.g. from "OB Candidates" to "Started Candidates"), keeping formulas working.

Usage:
    python3 move_candidate.py --xlsx data/OB_Candidates_Pay_Packages_Soloh.xlsx \
        --name "Alex Morgan Romero" --to "Started Candidates" [--out new.xlsx]

- The destination sheet is created (with the source sheet's header row) if missing.
- Formulas on the moved row, and on every row below the removed one, are renumbered
  so each row still points at itself. Absolute refs ($B$2, other sheets) are untouched.
- Excel recalculates all values when the file is next opened.
"""
import argparse, copy, re, sys
import openpyxl

REF = re.compile(r"(?<![A-Za-z$!'])(\$?)([A-Z]{1,3})(\d+)\b")

def renumber(formula, old, new):
    def sub(m):
        col_abs, col, row = m.groups()
        return f"{col_abs}{col}{new}" if int(row) == old else m.group(0)
    # leave references to other sheets alone ('Client Fees'!$B$2 etc.)
    parts = re.split(r"('[^']+'!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)", formula)
    return "".join(p if p.startswith("'") else REF.sub(sub, p) for p in parts)

def copy_row(src_ws, src_r, dst_ws, dst_r):
    for c in range(1, src_ws.max_column + 1):
        s = src_ws.cell(src_r, c); d = dst_ws.cell(dst_r, c)
        v = s.value
        d.value = renumber(v, src_r, dst_r) if isinstance(v, str) and v.startswith("=") else v
        if s.has_style:
            d.font = copy.copy(s.font); d.fill = copy.copy(s.fill); d.border = copy.copy(s.border)
            d.alignment = copy.copy(s.alignment); d.number_format = s.number_format

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True); ap.add_argument("--name", required=True)
    ap.add_argument("--to", required=True); ap.add_argument("--out")
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx)
    target = a.name.strip().lower()
    hit = None
    for ws in wb.worksheets:
        if ws.title == a.to or ws.cell(1, 1).value != "Candidate": continue
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v and str(v).strip().lower() == target:
                hit = (ws, r); break
        if hit: break
    if not hit: sys.exit(f"'{a.name}' not found on any candidate sheet.")
    src, row = hit
    if a.to in wb.sheetnames:
        dst = wb[a.to]
    else:
        dst = wb.create_sheet(a.to, index=wb.sheetnames.index(src.title) + 1)
        copy_row(src, 1, dst, 1)
        for k, dim in src.column_dimensions.items(): dst.column_dimensions[k].width = dim.width
        dst.freeze_panes = src.freeze_panes
    dst_row = next((r for r in range(2, dst.max_row + 2) if all(dst.cell(r, c).value in (None, "") for c in range(1, dst.max_column + 1))), dst.max_row + 1)
    copy_row(src, row, dst, dst_row)
    last = src.max_row
    for r in range(row, last):                      # shift rows below up by one, renumbering formulas
        copy_row(src, r + 1, src, r)
    for c in range(1, src.max_column + 1): src.cell(last, c).value = None
    wb.save(a.out or a.xlsx)
    print(f"Moved '{a.name}' from '{src.title}' row {row} to '{dst.title}' row {dst_row}.")

if __name__ == "__main__":
    main()
