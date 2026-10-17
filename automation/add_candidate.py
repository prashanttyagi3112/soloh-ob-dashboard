#!/usr/bin/env python3
"""
Add ONE approved candidate to the Margin Dashboard workbook (sheet "OB Candidates" by default).

    python3 automation/add_candidate.py --xlsx data/OB_Candidates_Pay_Packages_Soloh.xlsx --data new.json

new.json holds the values Prashant confirmed, keyed by dashboard field (see build.FIELD_ALIASES):
  candidate, client, facility, req, start (YYYY-MM-DD), end, shift, hrs_shift, guar_hrs,
  w2, ot, stipend, bill, team, recruiter, onetime, reason
- Formula columns (weeks, MSP fee, NBO, revenue, profit ...) are copied from the last row of the
  sheet and renumbered to the new row, so the workbook keeps calculating exactly as before.
- MSP % and NBO come from the formulas / Client Fees sheet unless given explicitly.
- Never guesses: any field not in new.json is left blank. Refuses a duplicate name.
"""
import argparse, copy, datetime as dt, json, os, sys
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from build import map_headers, norm              # same header matching as the dashboard build
from move_candidate import renumber

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True); ap.add_argument("--data", required=True)
    ap.add_argument("--sheet", default="OB Candidates"); ap.add_argument("--out")
    a = ap.parse_args()
    d = json.load(open(a.data))
    if not d.get("candidate"): sys.exit("candidate name is required")
    wb = openpyxl.load_workbook(a.xlsx)
    for ws in wb.worksheets:
        for r in range(2, ws.max_row + 1):
            if norm(ws.cell(r, 1).value) == norm(d["candidate"]) and ws.cell(r, 1).value:
                sys.exit(f"'{d['candidate']}' already exists on sheet '{ws.title}' row {r}. Nothing added.")
    if a.sheet not in wb.sheetnames: sys.exit(f"Sheet '{a.sheet}' not found. Sheets: {wb.sheetnames}")
    ws = wb[a.sheet]
    headers = [c.value for c in ws[1]]
    mapping, _ = map_headers(headers)
    last = max((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")), default=1)
    new = last + 1
    unknown = [k for k in d if k not in mapping]
    if unknown: sys.exit(f"No column on '{a.sheet}' for: {unknown}. Headers: {headers}")
    filled = []
    for ci in range(1, len(headers) + 1):
        src = ws.cell(last, ci) if last >= 2 else None
        dst = ws.cell(new, ci)
        field = next((f for f, i in mapping.items() if i == ci - 1), None)
        if field in d and d[field] not in (None, ""):
            v = d[field]
            if field in ("start", "end") and isinstance(v, str):
                v = dt.datetime.strptime(v[:10], "%Y-%m-%d")
            dst.value = v; filled.append(field)
        elif src is not None and isinstance(src.value, str) and src.value.startswith("="):
            dst.value = renumber(src.value, last, new); filled.append(f"{field or headers[ci-1]} (formula)")
        if src is not None and src.has_style:
            dst.font = copy.copy(src.font); dst.fill = copy.copy(src.fill); dst.border = copy.copy(src.border)
            dst.alignment = copy.copy(src.alignment); dst.number_format = src.number_format
    wb.save(a.out or a.xlsx)
    print(f"Added '{d['candidate']}' to '{a.sheet}' row {new}: {', '.join(filled)}")

if __name__ == "__main__":
    main()
